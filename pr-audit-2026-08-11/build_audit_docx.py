from __future__ import annotations

import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


GREEN = "173F35"
CORAL = "B95745"
CREAM = "FFF7EA"
GOLD = "C99B57"
INK = "263C48"
MUTED = "607080"
PALE = "F6F2EA"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=100, bottom=80, end=100) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def prevent_row_split(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run()
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = " PAGE "
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_char1, instr_text, fld_char2])


def add_hyperlink(paragraph, text: str, url: str, color: str = CORAL) -> None:
    # Keep the visible source label and URL in the portable review copy. Word's
    # PDF renderer can stall while resolving many mixed web/local relationships;
    # the native Google Doc is the linked delivery surface.
    run = paragraph.add_run(f"{text} ({url})")
    run.font.color.rgb = RGBColor.from_string(color)


INLINE_RE = re.compile(r"(\*\*[^*]+\*\*|`[^`]+`|\[[^\]]+\]\([^)]+\))")


def add_inline(paragraph, text: str, *, default_bold: bool = False) -> None:
    position = 0
    for match in INLINE_RE.finditer(text):
        if match.start() > position:
            run = paragraph.add_run(text[position : match.start()])
            run.bold = default_bold
        token = match.group(0)
        if token.startswith("**"):
            run = paragraph.add_run(token[2:-2])
            run.bold = True
        elif token.startswith("`"):
            run = paragraph.add_run(token[1:-1])
            run.font.name = "Consolas"
            run.font.size = Pt(8.5)
            run.font.color.rgb = RGBColor.from_string(MUTED)
        else:
            label, url = re.match(r"\[([^\]]+)\]\(([^)]+)\)", token).groups()
            add_hyperlink(paragraph, label, url)
        position = match.end()
    if position < len(text):
        run = paragraph.add_run(text[position:])
        run.bold = default_bold


def configure_styles(doc: Document) -> None:
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(9.4)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.08

    for style_name, size, color, before, after in (
        ("Title", 25, GREEN, 0, 12),
        ("Heading 1", 17, GREEN, 14, 6),
        ("Heading 2", 13, CORAL, 11, 4),
        ("Heading 3", 10.5, GREEN, 8, 3),
    ):
        style = styles[style_name]
        style.font.name = "Aptos Display"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    if "Audit Meta" not in styles:
        meta = styles.add_style("Audit Meta", WD_STYLE_TYPE.PARAGRAPH)
        meta.font.name = "Aptos"
        meta.font.size = Pt(9)
        meta.font.color.rgb = RGBColor.from_string(MUTED)
        meta.paragraph_format.space_after = Pt(2)

    if "Decision" not in styles:
        decision = styles.add_style("Decision", WD_STYLE_TYPE.PARAGRAPH)
        decision.font.name = "Aptos"
        decision.font.size = Pt(11)
        decision.font.bold = True
        decision.font.color.rgb = RGBColor.from_string(GREEN)
        decision.paragraph_format.space_before = Pt(7)
        decision.paragraph_format.space_after = Pt(8)


def parse_table(lines: list[str], index: int) -> tuple[list[list[str]], int]:
    rows: list[list[str]] = []
    while index < len(lines) and lines[index].strip().startswith("|"):
        row = [cell.strip() for cell in lines[index].strip().strip("|").split("|")]
        if not all(re.fullmatch(r":?-{3,}:?", cell or "") for cell in row):
            rows.append(row)
        index += 1
    return rows, index


def add_table(doc: Document, rows: list[list[str]]) -> None:
    if not rows:
        return
    width = max(len(row) for row in rows)
    table = doc.add_table(rows=len(rows), cols=width)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    for r_idx, row_values in enumerate(rows):
        row = table.rows[r_idx]
        # Allow long audit rows to continue across pages. Some rows contain
        # enough evidence text that forcing them to remain whole can create an
        # impossible pagination constraint in Word's PDF renderer.
        if r_idx == 0:
            repeat_table_header(row)
        for c_idx, cell in enumerate(row.cells):
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
            value = row_values[c_idx] if c_idx < len(row_values) else ""
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(1)
            paragraph.paragraph_format.line_spacing = 1.0
            add_inline(paragraph, value, default_bold=r_idx == 0)
            for run in paragraph.runs:
                run.font.size = Pt(7.4 if width >= 5 else 8.2)
                if r_idx == 0:
                    run.font.color.rgb = RGBColor(255, 255, 255)
            if r_idx == 0:
                set_cell_shading(cell, GREEN)
            elif r_idx % 2 == 0:
                set_cell_shading(cell, PALE)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)


def build(markdown_path: Path, output_path: Path) -> None:
    lines = markdown_path.read_text(encoding="utf-8").splitlines()
    doc = Document()
    configure_styles(doc)
    section = doc.sections[0]
    section.top_margin = Cm(1.55)
    section.bottom_margin = Cm(1.5)
    section.left_margin = Cm(1.65)
    section.right_margin = Cm(1.65)
    section.header_distance = Cm(0.65)
    section.footer_distance = Cm(0.65)

    header = section.header.paragraphs[0]
    header.text = "LOLO CARE  |  PR AGENCY WORK AUDIT"
    header.style = doc.styles["Audit Meta"]
    header.runs[0].font.bold = True
    header.runs[0].font.color.rgb = RGBColor.from_string(GREEN)
    header.runs[0].font.size = Pt(8)
    add_page_number(section.footer.paragraphs[0])

    title_seen = False
    index = 0
    while index < len(lines):
        raw = lines[index]
        stripped = raw.strip()
        if not stripped:
            index += 1
            continue

        if stripped.startswith("|") and index + 1 < len(lines) and lines[index + 1].strip().startswith("|"):
            rows, index = parse_table(lines, index)
            add_table(doc, rows)
            continue

        heading = re.match(r"^(#{1,3})\s+(.*)$", stripped)
        if heading:
            level = len(heading.group(1))
            text = heading.group(2)
            if level == 1 and not title_seen:
                p = doc.add_paragraph(style="Title")
                add_inline(p, text)
                p.paragraph_format.space_before = Pt(2)
                title_seen = True
            else:
                p = doc.add_paragraph(style=f"Heading {level}")
                add_inline(p, text)
            index += 1
            continue

        bullet = re.match(r"^-\s+(.*)$", stripped)
        numbered = re.match(r"^\d+\.\s+(.*)$", stripped)
        if bullet or numbered:
            p = doc.add_paragraph(style="List Bullet" if bullet else "List Number")
            p.paragraph_format.left_indent = Cm(0.45)
            p.paragraph_format.first_line_indent = Cm(-0.2)
            p.paragraph_format.space_after = Pt(2)
            add_inline(p, (bullet or numbered).group(1))
            index += 1
            continue

        p_style = "Audit Meta" if stripped.startswith("**Audit date:") or stripped.startswith("**Agency / owner") or stripped.startswith("**Source package:") or stripped.startswith("**Decision standard:") else None
        if stripped.startswith("**Overall score:") or stripped.startswith("**Recommendation:") or stripped.startswith("**Decision:"):
            p_style = "Decision"
        p = doc.add_paragraph(style=p_style)
        add_inline(p, stripped)
        index += 1

    props = doc.core_properties
    props.title = "LoLo Care PR Agency Work — Complete Quality and Fit Audit"
    props.subject = "Independent quality, evidence, execution-readiness, and visual audit"
    props.author = "Codex"
    props.keywords = "LoLo Care, PR, agency audit, media list, visibility, brand DNA"
    doc.save(output_path)


if __name__ == "__main__":
    if len(sys.argv) != 3:
        raise SystemExit("Usage: build_audit_docx.py INPUT.md OUTPUT.docx")
    build(Path(sys.argv[1]).resolve(), Path(sys.argv[2]).resolve())
